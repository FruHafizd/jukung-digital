"""Server aplikasi Jukung Digital: API JSON + menyajikan tampilan web dari folder web/."""
import glob
import io
import os
import subprocess
from datetime import date, datetime, timedelta
from pathlib import Path

from fastapi import Body, FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import db, demo, scraper, sumber

app = FastAPI(title="Jukung Digital")

STATUSES = ["baru", "dihubungi", "membalas", "tertarik", "deal", "ditolak", "tidak_merespons"]
TAHAP_PROYEK = ["materi", "desain", "revisi", "online", "selesai"]
STATUS_SELESAI = ("deal", "ditolak", "tidak_merespons")


@app.on_event("startup")
def _startup():
    scraper.mark_orphans()


@app.exception_handler(ValueError)
def _value_error(_: Request, exc: ValueError):
    return JSONResponse(status_code=400, content={"detail": str(exc)})


def _fokus():
    return db.get_setting("fokus_provinsi", "63")


def _pick(payload, allowed):
    return {k: v for k, v in payload.items() if k in allowed}


def _update(table, row_id, fields):
    if not fields:
        return
    sets = ", ".join(f"`{k}`=%s" for k in fields)
    db.execute(f"UPDATE {table} SET {sets} WHERE id=%s", (*[v if v != "" else None for v in fields.values()], row_id))


def _jsonable(rows):
    """Mengubah Decimal/date supaya bisa dikirim sebagai JSON."""
    def conv(v):
        if isinstance(v, (datetime, date)):
            return v.isoformat()
        if v.__class__.__name__ == "Decimal":
            return float(v)
        return v
    if isinstance(rows, dict):
        return {k: conv(v) for k, v in rows.items()}
    return [{k: conv(v) for k, v in r.items()} for r in rows]


# ---------------- Meta & wilayah ----------------

@app.get("/api/meta")
def meta():
    s = db.get_settings()
    fokus = s.get("fokus_provinsi", "63")
    return {
        "kategori": db.query("SELECT id, slug, nama, kata_kunci, tingkat, jam_kirim FROM kategori WHERE aktif=1 ORDER BY urutan"),
        "provinsi": db.query("SELECT kode, nama FROM wilayah WHERE tingkat=1 ORDER BY nama"),
        "fokus": db.query_one("SELECT kode, nama FROM wilayah WHERE kode=%s", (fokus,)),
        "paket": _jsonable(db.query("SELECT * FROM paket ORDER BY urutan")),
        "statuses": STATUSES,
        "pengaturan": {
            "nama_anda": s.get("nama_anda", ""), "brand": s.get("brand", "Jukung Digital"),
            "link_portofolio": s.get("link_portofolio", ""), "batas_wa_harian": int(s.get("batas_wa_harian") or 25),
            "biaya_perpanjangan": int(s.get("biaya_perpanjangan") or 0),
            "ada_api_key": bool(s.get(sumber.SUMBER[s.get("sumber_data", "serpapi")][1])),
            "sumber_label": sumber.LABEL[s.get("sumber_data", "serpapi")],
        },
        "ada_demo": bool(db.query_one("SELECT id FROM leads WHERE is_demo=1 LIMIT 1")),
    }


@app.get("/api/wilayah/{kode}/anak")
def wilayah_anak(kode: str):
    return db.query("SELECT kode, nama FROM wilayah WHERE induk=%s ORDER BY nama", (kode,))


# ---------------- Dashboard ----------------

@app.get("/api/dashboard")
def dashboard(scope: str = "fokus"):
    fokus = _fokus()
    where, params = ("WHERE l.provinsi_kode=%s", [fokus]) if scope == "fokus" else ("WHERE 1=1", [])
    counts = {r["status"]: r["n"] for r in db.query(
        f"SELECT status, COUNT(*) n FROM leads l {where} GROUP BY status", params)}
    total = sum(counts.values())
    c = lambda *ss: sum(counts.get(s, 0) for s in ss)  # noqa: E731
    corong = [
        {"tahap": "Lead terkumpul", "n": total},
        {"tahap": "Dihubungi", "n": total - c("baru")},
        {"tahap": "Membalas", "n": c("membalas", "tertarik", "deal", "ditolak")},
        {"tahap": "Tertarik", "n": c("tertarik", "deal")},
        {"tahap": "Deal", "n": c("deal")},
    ]
    agg = ("COUNT(*) total, SUM(l.status<>'baru') dihubungi, SUM(l.status IN ('membalas','tertarik','deal','ditolak')) membalas, "
           "SUM(l.status='deal') deal")
    # Alias "grup" (bukan "nama") supaya GROUP BY tidak tertukar dengan kolom leads.nama
    per_kategori = db.query(
        f"SELECT COALESCE(k.nama,'Kata kunci bebas') grup, {agg} FROM leads l LEFT JOIN kategori k ON k.id=l.kategori_id "
        f"{where} GROUP BY grup ORDER BY total DESC", params)
    group_col = "COALESCE(l.kota_nama,'Tidak diketahui')" if scope == "fokus" else "COALESCE(w.nama,'Tidak diketahui')"
    join = "" if scope == "fokus" else "LEFT JOIN wilayah w ON w.kode=l.provinsi_kode"
    per_wilayah = db.query(
        f"SELECT {group_col} grup, {agg} FROM leads l {join} {where} GROUP BY grup ORDER BY total DESC", params)
    for r in per_kategori + per_wilayah:
        r["nama"] = r.pop("grup")
    today = date.today()
    wa_hari_ini = db.query_one("SELECT COUNT(*) n FROM riwayat_kontak WHERE jenis='wa' AND DATE(created_at)=CURDATE()")["n"]
    followup = db.query_one(
        f"SELECT COUNT(*) n FROM leads l {where} AND l.followup_berikutnya <= CURDATE() "
        f"AND l.status NOT IN ('deal','ditolak','tidak_merespons')", params)["n"]
    bayar = db.query_one(
        "SELECT COALESCE(SUM(jumlah),0) n FROM pembayaran WHERE tanggal >= DATE_FORMAT(CURDATE(), '%%Y-%%m-01')")["n"]
    proyek_aktif = db.query_one("SELECT COUNT(*) n FROM proyek WHERE tahap<>'selesai'")["n"]
    tenggat = db.query(
        "SELECT id, nama_klien, deadline, tahap FROM proyek WHERE tahap NOT IN ('online','selesai') AND deadline IS NOT NULL "
        "AND deadline <= %s ORDER BY deadline LIMIT 6", (today + timedelta(days=14),))
    perpanjangan = db.query(
        "SELECT id, nama_klien, domain, tanggal_perpanjangan, biaya_perpanjangan FROM proyek "
        "WHERE tanggal_perpanjangan IS NOT NULL AND tanggal_perpanjangan <= %s ORDER BY tanggal_perpanjangan LIMIT 6",
        (today + timedelta(days=30),))
    aktivitas = db.query(
        "SELECT r.jenis, r.tahap, r.created_at, l.id lead_id, l.nama FROM riwayat_kontak r JOIN leads l ON l.id=r.lead_id "
        "ORDER BY r.id DESC LIMIT 8")
    return {
        "kpi": {
            "total": total, "belum": c("baru"), "followup": followup, "wa_hari_ini": wa_hari_ini,
            "batas_wa": int(db.get_setting("batas_wa_harian", "25")), "deal": c("deal"),
            "pendapatan_bulan_ini": int(bayar), "proyek_aktif": proyek_aktif,
        },
        "kuota": scraper.quota_status(),
        "corong": corong,
        "per_kategori": _jsonable(per_kategori),
        "per_wilayah": _jsonable(per_wilayah),
        "tenggat": _jsonable(tenggat),
        "perpanjangan": _jsonable(perpanjangan),
        "aktivitas": _jsonable(aktivitas),
    }


# ---------------- Scraping ----------------

@app.post("/api/scrape/estimasi")
def scrape_estimasi(params: dict = Body(...)):
    return scraper.estimate(params)


@app.post("/api/scrape")
def scrape_start(params: dict = Body(...)):
    return {"id": scraper.start_job(params)}


@app.get("/api/scrape/jobs")
def scrape_jobs():
    return _jsonable(db.query(
        "SELECT id, dibuat, selesai, status, ringkasan, total_pencarian, pencarian_selesai, request_dipakai, "
        "ditemukan, disimpan, pesan_error FROM riwayat_scraping ORDER BY id DESC LIMIT 20"))


@app.get("/api/scrape/jobs/{job_id}")
def scrape_job(job_id: int):
    row = db.query_one("SELECT * FROM riwayat_scraping WHERE id=%s", (job_id,))
    if not row:
        raise HTTPException(404, "Pencarian tidak ditemukan")
    row.pop("parameter", None)
    return _jsonable(row)


@app.post("/api/scrape/jobs/{job_id}/batal")
def scrape_cancel(job_id: int):
    scraper.cancel_job(job_id)
    return {"ok": True}


@app.get("/api/kuota")
def kuota():
    return scraper.quota_status()


# ---------------- Leads ----------------

SORTS = {
    "skor": "l.skor DESC, l.id DESC", "terbaru": "l.id DESC", "ulasan": "l.jumlah_ulasan DESC",
    "rating": "l.rating DESC", "nama": "l.nama ASC", "followup": "l.followup_berikutnya IS NULL, l.followup_berikutnya ASC",
}


def _lead_filters(q="", status="", kategori="", kabupaten="", website="", wa="", skor_min=0, scope="fokus", provinsi=""):
    where, params = ["1=1"], []
    if scope == "fokus":
        where.append("l.provinsi_kode=%s")
        params.append(_fokus())
    elif provinsi:
        where.append("l.provinsi_kode=%s")
        params.append(provinsi)
    if q:
        where.append("(l.nama LIKE %s OR l.alamat LIKE %s OR l.telepon LIKE %s)")
        params += [f"%{q}%"] * 3
    if status:
        where.append("l.status IN (" + ",".join(["%s"] * len(status.split(","))) + ")")
        params += status.split(",")
    if kategori:
        if kategori == "bebas":
            where.append("l.kategori_id IS NULL")
        else:
            where.append("l.kategori_id=%s")
            params.append(int(kategori))
    if kabupaten:
        where.append("l.kabupaten_kode=%s")
        params.append(kabupaten)
    if website:
        where.append("l.website_status=%s")
        params.append(website)
    if wa == "1":
        where.append("l.wa_nomor IS NOT NULL")
    if skor_min:
        where.append("l.skor >= %s")
        params.append(int(skor_min))
    return " AND ".join(where), params


LEAD_COLS = ("l.id, l.nama, l.kategori_id, k.nama kategori_nama, k.jam_kirim, l.kategori_google, l.alamat, l.provinsi_kode, "
             "l.kabupaten_kode, l.kota_nama, l.telepon, l.wa_nomor, l.website_status, l.website_url, l.sosmed_jenis, "
             "l.rating, l.jumlah_ulasan, l.maps_url, l.skor, l.status, l.tahap_followup, l.followup_berikutnya, l.catatan, "
             "l.is_demo, l.created_at, l.updated_at")


@app.get("/api/leads")
def leads(q: str = "", status: str = "", kategori: str = "", kabupaten: str = "", website: str = "", wa: str = "",
          skor_min: int = 0, scope: str = "fokus", provinsi: str = "", sort: str = "skor", page: int = 1, per_page: int = 50):
    where, params = _lead_filters(q, status, kategori, kabupaten, website, wa, skor_min, scope, provinsi)
    total = db.query_one(f"SELECT COUNT(*) n FROM leads l WHERE {where}", params)["n"]
    per_page = max(10, min(200, per_page))
    rows = db.query(
        f"SELECT {LEAD_COLS} FROM leads l LEFT JOIN kategori k ON k.id=l.kategori_id WHERE {where} "
        f"ORDER BY {SORTS.get(sort, SORTS['skor'])} LIMIT %s OFFSET %s",
        (*params, per_page, (max(1, page) - 1) * per_page))
    status_counts = {r["status"]: r["n"] for r in db.query(
        f"SELECT l.status, COUNT(*) n FROM leads l WHERE {_lead_filters(q, '', kategori, kabupaten, website, wa, skor_min, scope, provinsi)[0]} "
        "GROUP BY l.status", _lead_filters(q, '', kategori, kabupaten, website, wa, skor_min, scope, provinsi)[1])}
    return {"total": total, "items": _jsonable(rows), "status_counts": status_counts}


@app.get("/api/leads/{lead_id}")
def lead_detail(lead_id: int):
    row = db.query_one(f"SELECT {LEAD_COLS} FROM leads l LEFT JOIN kategori k ON k.id=l.kategori_id WHERE l.id=%s", (lead_id,))
    if not row:
        raise HTTPException(404, "Lead tidak ditemukan")
    row = _jsonable(row)
    row["riwayat"] = _jsonable(db.query("SELECT * FROM riwayat_kontak WHERE lead_id=%s ORDER BY id DESC", (lead_id,)))
    proyek = db.query_one("SELECT id FROM proyek WHERE lead_id=%s", (lead_id,))
    row["proyek_id"] = proyek["id"] if proyek else None
    return row


@app.patch("/api/leads/{lead_id}")
def lead_update(lead_id: int, payload: dict = Body(...)):
    fields = _pick(payload, {"status", "catatan", "followup_berikutnya"})
    if "status" in fields:
        if fields["status"] not in STATUSES:
            raise HTTPException(400, "Status tidak dikenal")
        old = db.query_one("SELECT status FROM leads WHERE id=%s", (lead_id,))
        if old and old["status"] != fields["status"]:
            db.execute("INSERT INTO riwayat_kontak (lead_id, jenis, tahap, isi) VALUES (%s,'status',%s,%s)",
                       (lead_id, fields["status"], f"Status: {old['status']} → {fields['status']}"))
        if fields["status"] in STATUS_SELESAI:
            fields["followup_berikutnya"] = None
    _update("leads", lead_id, fields)
    return lead_detail(lead_id)


@app.post("/api/leads/{lead_id}/kontak")
def lead_kontak(lead_id: int, payload: dict = Body(...)):
    """Mencatat pesan yang dikirim, lalu memajukan tahap follow-up otomatis (hari 1 → 3 → 7)."""
    lead = db.query_one("SELECT status, tahap_followup FROM leads WHERE id=%s", (lead_id,))
    if not lead:
        raise HTTPException(404, "Lead tidak ditemukan")
    tahap = payload.get("tahap", "umum")
    jenis = payload.get("jenis", "wa")
    db.execute("INSERT INTO riwayat_kontak (lead_id, jenis, tahap, isi) VALUES (%s,%s,%s,%s)",
               (lead_id, jenis, tahap, payload.get("isi", "")))
    today = date.today()
    status, stage, nxt = lead["status"], lead["tahap_followup"], None
    keep_next = True
    if tahap == "pembuka":
        stage, nxt = max(stage, 1), today + timedelta(days=2)
    elif tahap == "kirim_contoh":
        stage, nxt = max(stage, 2), today + timedelta(days=4)
    elif tahap == "follow_up":
        stage, nxt = 3, None
    elif tahap == "penawaran":
        nxt = today + timedelta(days=3)
    else:
        keep_next = False
    if status == "baru":
        status = "dihubungi"
    if status in STATUS_SELESAI:
        nxt, keep_next = None, True
    sql = "UPDATE leads SET status=%s, tahap_followup=%s" + (", followup_berikutnya=%s" if keep_next else "") + " WHERE id=%s"
    db.execute(sql, (status, stage, nxt, lead_id) if keep_next else (status, stage, lead_id))
    return lead_detail(lead_id)


@app.post("/api/leads/{lead_id}/deal")
def lead_deal(lead_id: int, payload: dict = Body(...)):
    lead = db.query_one("SELECT * FROM leads WHERE id=%s", (lead_id,))
    if not lead:
        raise HTTPException(404, "Lead tidak ditemukan")
    existing = db.query_one("SELECT id FROM proyek WHERE lead_id=%s", (lead_id,))
    if existing:
        return {"proyek_id": existing["id"]}
    lead_update(lead_id, {"status": "deal"})
    pid = _create_proyek({
        "lead_id": lead_id, "nama_klien": lead["nama"], "kontak_wa": lead["wa_nomor"],
        "paket_id": payload.get("paket_id"), "harga_deal": payload.get("harga_deal"),
        "deadline": payload.get("deadline"), "is_demo": lead["is_demo"],
    })
    return {"proyek_id": pid}


@app.delete("/api/leads/{lead_id}")
def lead_delete(lead_id: int):
    db.execute("DELETE FROM leads WHERE id=%s", (lead_id,))
    return {"ok": True}


@app.get("/api/followup")
def followup(scope: str = "fokus"):
    where, params = _lead_filters(scope=scope)
    base = (f"SELECT {LEAD_COLS} FROM leads l LEFT JOIN kategori k ON k.id=l.kategori_id WHERE {where} "
            "AND l.status NOT IN ('deal','ditolak','tidak_merespons') ")
    jatuh_tempo = db.query(base + "AND l.followup_berikutnya <= CURDATE() ORDER BY l.followup_berikutnya, l.skor DESC", params)
    mendatang = db.query(base + "AND l.followup_berikutnya > CURDATE() AND l.followup_berikutnya <= %s "
                         "ORDER BY l.followup_berikutnya, l.skor DESC", (*params, date.today() + timedelta(days=7)))
    belum = db.query(base + "AND l.status='baru' AND l.wa_nomor IS NOT NULL ORDER BY l.skor DESC LIMIT 15", params)
    habis = db.query(base + "AND l.tahap_followup=3 AND l.status='dihubungi' AND l.followup_berikutnya IS NULL "
                     "ORDER BY l.updated_at LIMIT 20", params)
    return {"jatuh_tempo": _jsonable(jatuh_tempo), "mendatang": _jsonable(mendatang),
            "belum_dihubungi": _jsonable(belum), "selesai_tanpa_balasan": _jsonable(habis)}


@app.get("/api/export/leads.xlsx")
def export_leads(q: str = "", status: str = "", kategori: str = "", kabupaten: str = "", website: str = "", wa: str = "",
                 skor_min: int = 0, scope: str = "fokus", provinsi: str = ""):
    where, params = _lead_filters(q, status, kategori, kabupaten, website, wa, skor_min, scope, provinsi)
    rows = db.query(f"SELECT {LEAD_COLS} FROM leads l LEFT JOIN kategori k ON k.id=l.kategori_id WHERE {where} "
                    "ORDER BY l.skor DESC", params)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lead UMKM"
    head = ["Nama usaha", "Kategori", "Kab/Kota", "Alamat", "Telepon", "Link WhatsApp", "Status website", "Medsos",
            "Rating", "Ulasan", "Skor", "Status", "Follow-up berikutnya", "Catatan", "Google Maps"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E5C57")
    for r in rows:
        ws.append([
            r["nama"], r["kategori_nama"] or r["kategori_google"], r["kota_nama"], r["alamat"], r["telepon"],
            f"https://wa.me/{r['wa_nomor']}" if r["wa_nomor"] else "",
            "Tidak ada" if r["website_status"] == "tidak_ada" else "Hanya medsos/marketplace", r["sosmed_jenis"] or "",
            float(r["rating"]) if r["rating"] is not None else None, r["jumlah_ulasan"], r["skor"],
            r["status"].replace("_", " "), r["followup_berikutnya"], r["catatan"] or "", r["maps_url"],
        ])
    for col, width in zip("ABCDEFGHIJKLMNO", [32, 22, 20, 45, 18, 30, 18, 14, 8, 8, 7, 14, 14, 30, 40]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"lead-umkm-{date.today():%Y%m%d}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ---------------- Template & jawaban ----------------

@app.get("/api/templates")
def templates():
    return db.query("SELECT * FROM template_pesan ORDER BY urutan, id")


@app.post("/api/templates")
def template_create(p: dict = Body(...)):
    tid = db.execute("INSERT INTO template_pesan (nama, tahap, wilayah, isi, urutan) VALUES (%s,%s,%s,%s,%s)",
                     (p.get("nama") or "Template baru", p.get("tahap", "umum"), p.get("wilayah", "semua"),
                      p.get("isi", ""), 99))
    return db.query_one("SELECT * FROM template_pesan WHERE id=%s", (tid,))


@app.put("/api/templates/{tid}")
def template_update(tid: int, p: dict = Body(...)):
    _update("template_pesan", tid, _pick(p, {"nama", "tahap", "wilayah", "isi"}))
    return db.query_one("SELECT * FROM template_pesan WHERE id=%s", (tid,))


@app.delete("/api/templates/{tid}")
def template_delete(tid: int):
    db.execute("DELETE FROM template_pesan WHERE id=%s", (tid,))
    return {"ok": True}


@app.get("/api/keberatan")
def keberatan():
    return db.query("SELECT * FROM jawaban_keberatan ORDER BY urutan, id")


@app.post("/api/keberatan")
def keberatan_create(p: dict = Body(...)):
    kid = db.execute("INSERT INTO jawaban_keberatan (judul, isi, urutan) VALUES (%s,%s,99)",
                     (p.get("judul") or "Keberatan baru", p.get("isi", "")))
    return db.query_one("SELECT * FROM jawaban_keberatan WHERE id=%s", (kid,))


@app.put("/api/keberatan/{kid}")
def keberatan_update(kid: int, p: dict = Body(...)):
    _update("jawaban_keberatan", kid, _pick(p, {"judul", "isi"}))
    return db.query_one("SELECT * FROM jawaban_keberatan WHERE id=%s", (kid,))


@app.delete("/api/keberatan/{kid}")
def keberatan_delete(kid: int):
    db.execute("DELETE FROM jawaban_keberatan WHERE id=%s", (kid,))
    return {"ok": True}


# ---------------- Paket ----------------

@app.put("/api/paket/{pid}")
def paket_update(pid: int, p: dict = Body(...)):
    _update("paket", pid, _pick(p, {"nama", "harga", "deskripsi", "fitur_default", "materi_default", "jatah_revisi"}))
    return _jsonable(db.query_one("SELECT * FROM paket WHERE id=%s", (pid,)))


# ---------------- Proyek ----------------

def _create_proyek(p):
    paket = db.query_one("SELECT * FROM paket WHERE id=%s", (p.get("paket_id"),)) if p.get("paket_id") else None
    harga = p.get("harga_deal") or (paket["harga"] if paket else 0)
    pid = db.execute(
        "INSERT INTO proyek (lead_id, nama_klien, kontak_wa, paket_id, harga_deal, deadline, jatah_revisi, "
        "biaya_perpanjangan, is_demo) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (p.get("lead_id"), p.get("nama_klien") or "Klien baru", p.get("kontak_wa"), paket["id"] if paket else None,
         int(harga or 0), p.get("deadline") or (date.today() + timedelta(days=14)),
         paket["jatah_revisi"] if paket else 3, int(db.get_setting("biaya_perpanjangan", "0") or 0),
         p.get("is_demo", 0)))
    if paket:
        rows = [(pid, "fitur", t.strip(), i) for i, t in enumerate((paket["fitur_default"] or "").splitlines()) if t.strip()]
        rows += [(pid, "materi", t.strip(), i) for i, t in enumerate((paket["materi_default"] or "").splitlines()) if t.strip()]
        with db.cursor() as cur:
            cur.executemany("INSERT INTO proyek_tugas (proyek_id, jenis, judul, urutan) VALUES (%s,%s,%s,%s)", rows)
    return pid


PROYEK_LIST = (
    "SELECT p.*, pk.nama paket_nama, "
    "(SELECT COUNT(*) FROM proyek_tugas t WHERE t.proyek_id=p.id AND t.jenis='fitur') fitur_total, "
    "(SELECT COUNT(*) FROM proyek_tugas t WHERE t.proyek_id=p.id AND t.jenis='fitur' AND t.selesai=1) fitur_selesai, "
    "(SELECT COUNT(*) FROM proyek_tugas t WHERE t.proyek_id=p.id AND t.jenis='materi') materi_total, "
    "(SELECT COUNT(*) FROM proyek_tugas t WHERE t.proyek_id=p.id AND t.jenis='materi' AND t.selesai=1) materi_selesai, "
    "(SELECT COALESCE(SUM(jumlah),0) FROM pembayaran b WHERE b.proyek_id=p.id AND b.jenis IN ('dp','pelunasan','lainnya')) dibayar, "
    "(SELECT COUNT(*) FROM revisi r WHERE r.proyek_id=p.id) revisi_terpakai "
    "FROM proyek p LEFT JOIN paket pk ON pk.id=p.paket_id"
)


@app.get("/api/proyek")
def proyek_list():
    return _jsonable(db.query(PROYEK_LIST + " ORDER BY p.urutan, p.id DESC"))


@app.post("/api/proyek")
def proyek_create(p: dict = Body(...)):
    return {"id": _create_proyek(p)}


@app.get("/api/proyek/{pid}")
def proyek_detail(pid: int):
    row = db.query_one(PROYEK_LIST + " WHERE p.id=%s", (pid,))
    if not row:
        raise HTTPException(404, "Proyek tidak ditemukan")
    row = _jsonable(row)
    row["tugas"] = db.query("SELECT * FROM proyek_tugas WHERE proyek_id=%s ORDER BY jenis, urutan, id", (pid,))
    row["pembayaran"] = _jsonable(db.query("SELECT * FROM pembayaran WHERE proyek_id=%s ORDER BY tanggal, id", (pid,)))
    row["revisi"] = _jsonable(db.query("SELECT * FROM revisi WHERE proyek_id=%s ORDER BY id", (pid,)))
    row["catatan"] = _jsonable(db.query("SELECT * FROM catatan WHERE proyek_id=%s ORDER BY disematkan DESC, updated_at DESC", (pid,)))
    if row["lead_id"]:
        row["lead"] = _jsonable(db.query_one("SELECT id, alamat, telepon, maps_url, kota_nama FROM leads WHERE id=%s",
                                             (row["lead_id"],)) or {})
    return row


@app.patch("/api/proyek/{pid}")
def proyek_update(pid: int, p: dict = Body(...)):
    fields = _pick(p, {"nama_klien", "kontak_wa", "paket_id", "harga_deal", "tahap", "urutan", "deadline", "domain",
                       "tanggal_online", "tanggal_perpanjangan", "biaya_perpanjangan", "jatah_revisi"})
    if fields.get("tahap") == "online":
        cur = db.query_one("SELECT tanggal_online, tanggal_perpanjangan FROM proyek WHERE id=%s", (pid,))
        if cur and not cur["tanggal_online"]:
            fields["tanggal_online"] = date.today()
            if not cur["tanggal_perpanjangan"]:
                fields["tanggal_perpanjangan"] = date.today() + timedelta(days=365)
    _update("proyek", pid, fields)
    return proyek_detail(pid)


@app.post("/api/proyek/urutan")
def proyek_urutan(p: dict = Body(...)):
    """Menyimpan posisi kartu di papan: {tahap: [id, id, ...]}"""
    for tahap, ids in p.items():
        if tahap in TAHAP_PROYEK:
            for i, pid in enumerate(ids):
                db.execute("UPDATE proyek SET tahap=%s, urutan=%s WHERE id=%s", (tahap, i, int(pid)))
    return {"ok": True}


@app.delete("/api/proyek/{pid}")
def proyek_delete(pid: int):
    db.execute("DELETE FROM proyek WHERE id=%s", (pid,))
    return {"ok": True}


@app.post("/api/proyek/{pid}/tugas")
def tugas_create(pid: int, p: dict = Body(...)):
    tid = db.execute("INSERT INTO proyek_tugas (proyek_id, jenis, judul, urutan) VALUES (%s,%s,%s,99)",
                     (pid, p.get("jenis", "fitur"), p.get("judul", "").strip() or "Tugas baru"))
    return db.query_one("SELECT * FROM proyek_tugas WHERE id=%s", (tid,))


@app.patch("/api/tugas/{tid}")
def tugas_update(tid: int, p: dict = Body(...)):
    _update("proyek_tugas", tid, _pick(p, {"judul", "selesai"}))
    return db.query_one("SELECT * FROM proyek_tugas WHERE id=%s", (tid,))


@app.delete("/api/tugas/{tid}")
def tugas_delete(tid: int):
    db.execute("DELETE FROM proyek_tugas WHERE id=%s", (tid,))
    return {"ok": True}


@app.post("/api/proyek/{pid}/pembayaran")
def bayar_create(pid: int, p: dict = Body(...)):
    db.execute("INSERT INTO pembayaran (proyek_id, jenis, jumlah, tanggal, catatan) VALUES (%s,%s,%s,%s,%s)",
               (pid, p.get("jenis", "dp"), int(p.get("jumlah") or 0), p.get("tanggal") or date.today(), p.get("catatan")))
    if p.get("jenis") == "perpanjangan":
        db.execute("UPDATE proyek SET tanggal_perpanjangan = DATE_ADD(COALESCE(tanggal_perpanjangan, CURDATE()), "
                   "INTERVAL 1 YEAR) WHERE id=%s", (pid,))
    return proyek_detail(pid)


@app.delete("/api/pembayaran/{bid}")
def bayar_delete(bid: int):
    db.execute("DELETE FROM pembayaran WHERE id=%s", (bid,))
    return {"ok": True}


@app.post("/api/proyek/{pid}/revisi")
def revisi_create(pid: int, p: dict = Body(...)):
    db.execute("INSERT INTO revisi (proyek_id, isi) VALUES (%s,%s)", (pid, p.get("isi", "").strip() or "Revisi"))
    return proyek_detail(pid)


@app.patch("/api/revisi/{rid}")
def revisi_update(rid: int, p: dict = Body(...)):
    _update("revisi", rid, _pick(p, {"isi", "selesai"}))
    return {"ok": True}


@app.delete("/api/revisi/{rid}")
def revisi_delete(rid: int):
    db.execute("DELETE FROM revisi WHERE id=%s", (rid,))
    return {"ok": True}


# ---------------- Catatan ----------------

@app.get("/api/catatan")
def catatan_list(proyek_id: int | None = None, q: str = ""):
    where, params = (["proyek_id=%s"], [proyek_id]) if proyek_id else (["proyek_id IS NULL"], [])
    if q:
        where.append("(judul LIKE %s OR isi LIKE %s)")
        params += [f"%{q}%", f"%{q}%"]
    return _jsonable(db.query(f"SELECT * FROM catatan WHERE {' AND '.join(where)} ORDER BY disematkan DESC, updated_at DESC",
                              params))


@app.post("/api/catatan")
def catatan_create(p: dict = Body(...)):
    cid = db.execute("INSERT INTO catatan (proyek_id, judul, isi) VALUES (%s,%s,%s)",
                     (p.get("proyek_id"), p.get("judul") or "Tanpa judul", p.get("isi", "")))
    return _jsonable(db.query_one("SELECT * FROM catatan WHERE id=%s", (cid,)))


@app.put("/api/catatan/{cid}")
def catatan_update(cid: int, p: dict = Body(...)):
    _update("catatan", cid, _pick(p, {"judul", "isi", "disematkan"}))
    return _jsonable(db.query_one("SELECT * FROM catatan WHERE id=%s", (cid,)))


@app.delete("/api/catatan/{cid}")
def catatan_delete(cid: int):
    db.execute("DELETE FROM catatan WHERE id=%s", (cid,))
    return {"ok": True}


# ---------------- Pengaturan ----------------

KUNCI_RAHASIA = {"api_key", "serpapi_key", "apify_token"}
PENGATURAN_KEYS = KUNCI_RAHASIA | {
    "nama_anda", "brand", "link_portofolio", "fokus_provinsi", "batas_bulanan", "batas_harian_api",
    "alokasi_fokus_persen", "batas_wa_harian", "biaya_perpanjangan", "sumber_data", "cadangan_apify",
    "batas_bulanan_serpapi", "batas_bulanan_apify",
}


def _samarkan(key):
    return (key[:6] + "…" + key[-4:]) if len(key) > 12 else ("" if not key else "••••")


@app.get("/api/pengaturan")
def pengaturan():
    s = {k: v for k, v in db.get_settings().items() if k in PENGATURAN_KEYS}
    for k in KUNCI_RAHASIA:
        s["ada_" + k] = bool(s.get(k))
        s[k] = _samarkan(s.get(k) or "")
    s["kuota"] = {k: scraper.quota_status(k) for k in sumber.SUMBER}
    return s


@app.put("/api/pengaturan")
def pengaturan_update(p: dict = Body(...)):
    for k, v in p.items():
        if k not in PENGATURAN_KEYS:
            continue
        if k in KUNCI_RAHASIA and (not v or "…" in v or v == "••••"):
            continue  # nilai tersamar = tidak diubah
        if k == "sumber_data" and v not in sumber.SUMBER:
            continue
        db.set_setting(k, str(v).strip())
    return pengaturan()


@app.post("/api/pengaturan/hapus-kunci/{nama}")
def hapus_kunci(nama: str):
    if nama in KUNCI_RAHASIA:
        db.set_setting(nama, "")
    return pengaturan()


@app.post("/api/pengaturan/tes/{kode}")
def tes_sumber(kode: str):
    if kode not in sumber.SUMBER:
        raise HTTPException(404, "Sumber tidak dikenal")
    try:
        pesan = scraper.tes_sumber(kode)
    except (scraper.ApiError, scraper.QuotaStop, ValueError) as e:
        raise HTTPException(400, str(e)) from e
    return {"ok": True, "pesan": pesan}


def _find_mysqldump():
    for pattern in [r"C:\laragon\bin\mysql\*\bin\mysqldump.exe", r"D:\laragon\bin\mysql\*\bin\mysqldump.exe",
                    r"E:\laragon\bin\mysql\*\bin\mysqldump.exe"]:
        found = sorted(glob.glob(pattern))
        if found:
            return found[-1]
    return "mysqldump"


@app.post("/api/backup")
def backup():
    s = db.db_settings()
    out_dir = db.ROOT / "backups"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / f"jukung_{datetime.now():%Y%m%d_%H%M%S}.sql"
    env = {**os.environ, "MYSQL_PWD": s["password"]}
    cmd = [_find_mysqldump(), f"--host={s['host']}", f"--port={s['port']}", f"--user={s['user']}",
           "--single-transaction", "--default-character-set=utf8mb4", f"--result-file={out}", s["name"]]
    try:
        r = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=300)
    except FileNotFoundError as e:
        raise HTTPException(500, "mysqldump tidak ditemukan. Pastikan Laragon terpasang.") from e
    if r.returncode != 0:
        raise HTTPException(500, "Backup gagal: " + r.stderr.strip()[:300])
    return {"ok": True, "file": str(out), "ukuran_kb": round(out.stat().st_size / 1024)}


# ---------------- Data contoh ----------------

@app.post("/api/demo")
def demo_isi():
    return {"dibuat": demo.isi()}


@app.delete("/api/demo")
def demo_hapus():
    demo.hapus()
    return {"ok": True}


# ---------------- Tampilan web ----------------

WEB = Path(__file__).resolve().parent.parent / "web"
app.mount("/", StaticFiles(directory=WEB, html=True), name="web")
